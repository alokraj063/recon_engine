"""
Oracle "External AR Statement" (the aging report Finance sends monthly)
-> one row per open invoice, source-native column names (silver).

The workbook carries several sheets (Aging, Data, Pivot, ...); the
invoice list is the sheet whose header row names REQUIRED_HEADERS, found
by TEXT in the first rows of every sheet, never by position. The header
spans two rows — the second only splits a few amount/aging columns into
"Original amount" / "Open amount" — so the first row's text is the
column name, suffixed with the second row's where a name repeats
("Current", "Past due"). Blank rows and a total row with no invoice
number are skipped.

The statement date is the "AR Statement <Mon YYYY>" title above the
header (-> the last day of that month); the adapter falls back to its
params or the newest invoice date when a file has no such title.
"""

import calendar
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd

from .bill_status import _XlsWorkbook, _norm_header

REQUIRED_HEADERS = ("Invoice number", "Due date", "Functional amount - open")
MAX_HEADER_SCAN_ROWS = 15
_TITLE_RE = re.compile(r"AR\s+Statement\s+([A-Za-z]{3,9})\s+(\d{4})", re.I)


class ArStatementFormatError(ValueError):
    pass


def _open(path):
    if Path(path).suffix.lower() == ".xls":
        return _XlsWorkbook(path)
    # read_only: the Data sheet runs to ~18k rows x 39 columns
    return openpyxl.load_workbook(path, data_only=True, read_only=True)


def _title_date(rows) -> date | None:
    for row in rows:
        for v in row:
            m = _TITLE_RE.search(str(v)) if isinstance(v, str) else None
            if m:
                try:
                    month = datetime.strptime(m.group(1)[:3].title(), "%b").month
                except ValueError:
                    continue
                year = int(m.group(2))
                return date(year, month, calendar.monthrange(year, month)[1])
    return None


def _columns(head, sub):
    """Unique column names from the two header rows."""
    names, seen = [], {}
    for i, h in enumerate(head):
        name = re.sub(r"\s+", " ", str(h).replace("\xa0", " ")).strip() if h is not None else ""
        if not name:
            names.append(None)
            continue
        extra = sub[i] if sub is not None and i < len(sub) else None
        base = name
        if base in seen:
            name = f"{base} ({extra})" if extra else f"{base} #{seen[base] + 1}"
        seen[base] = seen.get(base, 0) + 1
        names.append(name)
    return names


def parse_ar_statement(path):
    """Returns (DataFrame, meta). meta: sheet, header_row, statement_date."""
    wb = _open(path)
    required = {_norm_header(h) for h in REQUIRED_HEADERS}
    for name in wb.sheetnames:
        ws = wb[name]
        top = list(ws.iter_rows(min_row=1, max_row=MAX_HEADER_SCAN_ROWS,
                                values_only=True))
        for i, row in enumerate(top):
            if not required <= {_norm_header(v) for v in row}:
                continue
            sub = top[i + 1] if i + 1 < len(top) else None
            cols = _columns(row, sub)
            # the second header row is labels, not data, when it carries
            # no invoice number
            inv = [_norm_header(c) for c in cols].index("invoice number")
            first_data = i + 2 if sub is not None and sub[inv] is None else i + 1
            records = []
            for r in ws.iter_rows(min_row=first_data + 1, values_only=True):
                if inv >= len(r) or r[inv] is None or not str(r[inv]).strip():
                    continue
                records.append({c: r[j] if j < len(r) else None
                                for j, c in enumerate(cols) if c})
            df = pd.DataFrame(records, columns=[c for c in cols if c])
            df["Invoice number"] = df["Invoice number"].map(
                lambda v: str(int(v)) if isinstance(v, float) and v.is_integer()
                else str(v).strip())
            meta = {"sheet": name, "header_row": i + 1,
                    "statement_date": _title_date(top[:i])}
            return df, meta
    raise ArStatementFormatError(
        f"no sheet has an AR invoice header ({', '.join(REQUIRED_HEADERS)})")
