"""
The Daily Collection export: 4-4-5 fiscal calendar, the Oracle AR
statement source, and db/daily_collection_export.py + its routes —
confirmed matches -> the collections team's month-to-date sheet, one row
per (credit, bill), blanks + issues where the data cannot say.
"""

import sys
import uuid
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, select

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from db import SessionLocal, init_db  # noqa: E402
from db import collection as dc  # noqa: E402
from db import daily_collection_export as x  # noqa: E402
from db.bronze import register_file  # noqa: E402
from db.models import (AuditLog, Base, BronzeFile, Customer, GoldArInvoice,  # noqa: E402
                       GoldBankTxn, GoldBill, MatchLedger, MatchLedgerBill,
                       MatchRuleSetRow)
from db.storage import storage  # noqa: E402
from recon.fiscal import FiscalCalendar, FiscalCalendarError  # noqa: E402
from recon.sources import BY_KEY  # noqa: E402


# --- fiscal calendar ------------------------------------------------------

@pytest.mark.parametrize("day,month,week", [
    # boundaries read off the team's own Daily Collection sheets
    ("2026-06-29", "Jul 26", "1st Week"),
    ("2026-07-26", "Jul 26", "4th Week"),
    ("2026-07-27", "Aug 26", "1st Week"),
    ("2026-08-23", "Aug 26", "4th Week"),
    ("2026-08-24", "Sep 26", "1st Week"),
    ("2026-08-31", "Sep 26", "2nd Week"),
    ("2026-09-21", "Sep 26", "5th Week"),
    ("2026-09-27", "Sep 26", "5th Week"),
    ("2026-09-28", "Oct 26", "1st Week"),
    ("2025-12-29", "Jan 26", "1st Week"),
    ("2025-12-28", "Dec 25", "5th Week"),
])
def test_fiscal_periods(day, month, week):
    p = FiscalCalendar().period(day)
    assert (p.month_label, p.week_label) == (month, week)


def test_fiscal_53_week_year_and_overrides():
    cal = FiscalCalendar()
    # 2031 ends on Sun 28 Dec 2031 after starting Mon 30 Dec 2030: 52 weeks;
    # 2032 (last Sunday = 26 Dec 2032) — every year is whole weeks
    for y in range(2024, 2040):
        days = (cal.year_end(y) - cal.year_start(y)).days + 1
        assert days in (364, 371) and cal.year_start(y).weekday() == 0
        bounds = cal.month_bounds(y)
        assert bounds[0][0] == cal.year_start(y) and bounds[-1][1] == cal.year_end(y)
    moved = FiscalCalendar(year_starts={"2027": "2027-01-04"})
    assert moved.period("2027-01-01").month_label == "Dec 26"
    with pytest.raises(FiscalCalendarError):
        FiscalCalendar(year_starts={"2027": "2027-01-05"})      # a Tuesday
    with pytest.raises(FiscalCalendarError):
        FiscalCalendar(pattern=[4] * 12)                          # 48 weeks


# --- the AR statement source ----------------------------------------------

def _ar_workbook(path, title="AR Statement Aug 2026", rows=()):
    wb = Workbook()
    ws = wb.active
    ws.title = "Aging"
    ws["A1"] = "not this sheet"
    data = wb.create_sheet("Data")
    data["A1"] = title
    head = ["Category", "Sub category", "Invoice number", "Customer name",
            "Sales Order Type", "Sales Rep", "Invoice date", "Due date",
            "Functional Currency", "Functional amount", "Functional amount - open",
            "Current", "Current", "Operating unit"]
    sub = [None] * 9 + ["Original amount", "Open amount", "< 30 Days", "> 30 Days", None]
    data.append([])
    data.append(head)
    data.append(sub)
    for r in rows:
        data.append(r)
    data.append([None, None, None, "Total", None, None, None, None, None, 1, 1])
    wb.save(path)
    return path


def test_ar_adapter(tmp_path):
    p = _ar_workbook(tmp_path / "ar.xlsx", rows=[
        ["CS", "CS-Others", 1331000016, "SR", None, "WABHSR CS -Spares Sales",
         datetime(2026, 7, 1), datetime(2026, 8, 20), "INR", 1000.0, 900.0, None, 5, "Hosur B&S"],
    ])
    a = BY_KEY["oracle_ar"]
    silver = a.parse(p, {})
    assert silver.meta["sheet"] == "Data" and silver.meta["statement_date"] == "2026-08-31"
    assert "Current (> 30 Days)" in silver.frames["ar_statement"].columns
    gold = a.to_gold(silver, {})["ar_invoices"]
    assert len(gold) == 1                                       # total row skipped
    r = gold.iloc[0]
    assert r["invoice_number"] == "1331000016"
    assert r["functional_amount_open"] == 900.0
    assert r["statement_date"].date() == date(2026, 8, 31)
    assert r["due_date"].date() == date(2026, 8, 20)
    a.selfcheck({"ar_invoices": gold}, p, {})


# --- the export -----------------------------------------------------------

def _wipe(pk, key):
    with SessionLocal() as s:
        ids = [m.id for m in s.query(MatchLedger).filter_by(customer_id=pk)]
        if ids:
            s.execute(delete(MatchLedgerBill).where(MatchLedgerBill.match_ledger_id.in_(ids)))
        tables = [t for t in reversed(Base.metadata.sorted_tables)
                  if "customer_id" in t.c and t.name != "customers"]
        for t in tables:
            s.execute(delete(t).where(t.c.customer_id == pk))
        s.execute(delete(Customer).where(Customer.id == pk))
        s.commit()
    import shutil
    shutil.rmtree(storage.root / "bronze" / key, ignore_errors=True)


@pytest.fixture()
def world(tmp_path):
    init_db()
    key = f"dc-{uuid.uuid4().hex[:8]}"
    with SessionLocal() as s:
        c = Customer(key=key, name="daily collection test")
        s.add(c)
        s.flush()
        p = tmp_path / "f.txt"
        p.write_text(uuid.uuid4().hex)
        fid = register_file(s, c, "bill_status", p, p.name).id
        seq = iter(range(1000))

        def bill(number, zone, net, bill_date=date(2026, 6, 1)):
            b = GoldBill(customer_id=c.id, bronze_file_id=fid, row_seq=next(seq),
                         bill_number=number, submission_ref=f"CO6-{number}",
                         zone=zone, operating_unit="Hosur", net_payable_amount=net,
                         bill_date=bill_date, bill_status="PAYMENT MADE")
            s.add(b)
            s.flush()
            return b

        def match(ref, amount, bills, day, status="LOCKED", zone=None):
            t = GoldBankTxn(customer_id=c.id, bronze_file_id=fid, row_seq=next(seq),
                            bank_ref=ref, amount=amount, value_date=day,
                            zone_guess=zone, used_in_recon=True)
            s.add(t)
            s.flush()
            m = MatchLedger(customer_id=c.id, match_id="m", gold_bank_txn_id=t.id,
                            confidence="HIGH", status=status, seq=next(seq))
            s.add(m)
            s.flush()
            for b in bills:
                s.add(MatchLedgerBill(match_ledger_id=m.id, gold_bill_id=b.id,
                                      role="picked"))

        def ar(number, statement, due, open_amt, rep="WABHSR CS -Spares Sales",
               order_type=None, category=None, subcategory=None):
            s.add(GoldArInvoice(customer_id=c.id, bronze_file_id=fid, row_seq=next(seq),
                                statement_date=statement, invoice_number=number,
                                due_date=due, sales_rep=rep, sales_order_type=order_type,
                                category=category, subcategory=subcategory,
                                functional_amount=open_amt, functional_amount_open=open_amt))

        # Mon 24 Aug 2026 (Sep 26, 1st week): one bill, due 1 Sep -> inside the month
        match("UTR-A", 1000.0, [bill("1331000001", "SR", 1000.0)], date(2026, 8, 24))
        ar("1331000001", date(2026, 7, 31), date(2026, 9, 1), 1100.0)
        ar("1331000001", date(2026, 8, 31), date(2026, 9, 1), 0.0)
        # same day: a batch of two bills; the first is overdue by then (OD)
        # and Friction; the second is in no AR statement
        match("UTR-B", 500.0, [bill("5110000812", "ECOR", 300.0),
                               bill("9999999999", "ECOR", 200.0)], date(2026, 8, 24))
        ar("5110000812", date(2026, 7, 31), date(2026, 8, 1), 320.0, rep="WABINF Break Pad")
        # Mon 31 Aug (2nd week): a rep the master lacks, Finance's own label
        # used; an unknown zone
        match("UTR-C", 70.0, [bill("1331000003", "MTPK", 70.0)], date(2026, 8, 31))
        ar("1331000003", date(2026, 8, 31), date(2026, 12, 1), 70.0,
           rep="WABHSR MP - Metro Products", category="METRO", subcategory="Metro")
        # next fiscal month (Oct 26): its own sheet, its own totals
        match("UTR-D", 40.0, [bill("1331000004", "CR", 40.0)], date(2026, 9, 28))
        # excluded: an OE (production unit) bill — the sheet is TSG only
        match("UTR-OE", 90.0, [bill("1331000099", "ICF", 90.0)], date(2026, 8, 24))
        # excluded: an OPEN review match
        match("UTR-OPEN", 10.0, [bill("X-1", "CR", 10.0)], date(2026, 8, 24), status="OPEN")
        s.commit()
        pk = c.id
    yield {"pk": pk, "key": key}
    _wipe(pk, key)


def test_rows(world):
    with SessionLocal() as s:
        rows = x.build_rows(s, world["pk"], date(2026, 8, 24), date(2026, 9, 30))
    assert [r["Bank Ref"] for r in rows] == ["UTR-A", "UTR-B", "UTR-B", "UTR-C", "UTR-D"]
    a, b1, b2, c, d = rows

    # credit fields on the credit's first row only; SL.NO counts credits
    assert (a["SL.NO"], b1["SL.NO"], b2["SL.NO"], c["SL.NO"]) == (1, 2, None, 3)
    assert b1["EFT AMOUNT"] == 500.0 and b2["EFT AMOUNT"] is None
    assert d["SL.NO"] == 1                                  # a new sheet restarts

    # zone directory -> Region / RLY
    assert (a["Region"], a["RLY"]) == ("SOUTH", "SR")
    assert (b1["Region"], b1["RLY"]) == ("EAST", "ECOR")

    # AR: the statement known that day (31 Jul), not the later one
    assert a["Invoice Value"] == 1100.0 and a["ar_statement_date"] == date(2026, 7, 31)
    assert a["OD/NOD"] == "OD Sep 26" and b1["OD/NOD"] == "OD" and c["OD/NOD"] == "NOD"
    assert (a["Category"], a["BRANCH"]) == ("S&T", "ST")
    assert (b1["Category"], b1["BRANCH"]) == ("Rohtak Friction", "FR")
    assert b2["issues"] == ["NO_AR_INVOICE"] and b2["Invoice Value"] is None
    # a rep the master lacks falls back to the statement's own labels;
    # Metro has no branch code yet
    assert (c["Category"], c["BRANCH"]) == ("METRO", None)
    assert set(c["issues"]) == {"NO_ZONE", "NO_BRANCH"}

    # the day's collection and the month to date sit on the day's last row
    assert [r["DAILY COLLECTION"] for r in rows] == [None, None, 1500.0, 70.0, 40.0]
    assert [r["TOTAL COLLECTION"] for r in rows] == [None, None, 1500.0, 1570.0, 40.0]
    assert [r["Week"] for r in rows] == ["1st Week"] * 3 + ["2nd Week", "1st Week"]
    assert [r["month"] for r in rows] == ["Sep 26"] * 4 + ["Oct 26"]

    summ = x.summarize(rows)
    assert summ["credits"] == 4 and summ["eft_amount"] == 1610.0
    assert summ["by_month"] == {"Sep 26": 4, "Oct 26": 1}


def test_od_nod_is_a_fiscal_month_due_date_bucket():
    from recon.fiscal import FiscalCalendar
    cal, credit = FiscalCalendar(), date(2026, 9, 10)     # fiscal Sep 26: 24 Aug-27 Sep
    assert x.od_nod(cal, credit, date(2026, 8, 23)) == "OD"            # before the month
    assert x.od_nod(cal, credit, date(2026, 8, 24)) == "OD Sep 26"     # first day
    assert x.od_nod(cal, credit, date(2026, 9, 27)) == "OD Sep 26"     # last day
    assert x.od_nod(cal, credit, date(2026, 9, 28)) == "OD Oct 26"     # next month
    assert x.od_nod(cal, credit, date(2026, 10, 26)) == "NOD"
    assert x.od_nod(cal, credit, None) is None
    # due AFTER the credit but inside the month is still "OD <month>"
    assert x.od_nod(cal, date(2026, 9, 1), date(2026, 9, 20)) == "OD Sep 26"


def test_window_snaps_to_fiscal_month(world):
    with SessionLocal() as s:
        assert x.window(s, world["pk"], date(2026, 9, 3), date(2026, 9, 3)) == \
            (date(2026, 8, 24), date(2026, 9, 3))
        # no dates: the newest confirmed credit's month, up to that day
        assert x.window(s, world["pk"], None, None) == (date(2026, 9, 28), date(2026, 9, 28))


def test_download(world):
    from app.main import app
    with TestClient(app) as client:
        res = client.get("/api/export/daily-collection",
                         params={"customer_id": world["key"], "from": "2026-08-24",
                                 "to": "2026-09-30"})
        assert res.status_code == 200
        assert "DAILY%20COLLECTION%2030%20Sep%2026.xlsx" in res.headers["content-disposition"]
        wb = load_workbook(BytesIO(res.content))
        assert wb.sheetnames == ["Sep 26", "Oct 26"]
        ws = wb["Sep 26"]
        assert [ws.cell(row=4, column=i).value for i in (1, 2, 4, 13, 14, 18, 19)] == [
            "SL.NO", "EFT DATE", "Bank Ref", "DAILY COLLECTION", "TOTAL COLLECTION",
            None, "Week"]
        assert ws["D5"].value == "UTR-A" and ws["H5"].value == 1331000001
        assert ws["M7"].value == "=SUM(J5:J7)" and ws["N7"].value == "=M7"
        assert ws["M8"].value == "=SUM(J8:J8)" and ws["N8"].value == "=M8+N7"
        assert ws["G3"].value == "=SUM(G5:G8)" and ws["R5"].value == "S&T"

        prev = client.get("/api/export/daily-collection/preview",
                          params={"customer_id": world["key"], "from": "2026-09-01",
                                  "to": "2026-09-01"}).json()
        assert prev["from"] == "2026-08-24" and prev["summary"]["credits"] == 3
        assert {a["statement_date"] for a in prev["ar_statements"]} == {"2026-07-31", "2026-08-31"}
    with SessionLocal() as s:
        ev = s.query(AuditLog).filter_by(
            customer_id=world["pk"], event_type="export.daily_collection_downloaded").one()
        assert ev.details["rows"] == 5 and "UTR-A" not in str(ev.details)


def test_collection_settings(world):
    from app.main import app
    with TestClient(app) as client:
        url = f"/api/customers/{world['key']}/collection"
        got = client.get(url).json()
        assert got["branch_codes"] == dc.DEFAULT_BRANCH_CODES
        assert set(got["defaults"]) == {"fiscal_calendar", "category_master",
                                        "branch_codes", "recipients", "segments"}
        body = {**{k: got[k] for k in ("fiscal_calendar", "category_master")},
                "branch_codes": {**got["branch_codes"], "Metro": "metro"},
                "recipients": ["a@x.com", "A@x.com", " "]}
        put = client.put(url, json=body)
        assert put.status_code == 200, put.text
        assert put.json()["branch_codes"]["Metro"] == "METRO"
        assert put.json()["recipients"] == ["a@x.com"]
        assert put.json()["defaults"] == ["category_master", "fiscal_calendar", "segments"]
        # the new code reaches the export at once
        prev = client.get("/api/export/daily-collection/preview",
                          params={"customer_id": world["key"], "from": "2026-08-31",
                                  "to": "2026-08-31"}).json()
        assert [r["BRANCH"] for r in prev["rows"] if r["Bank Ref"] == "UTR-C"] == ["METRO"]

        # OE too: the production-unit credit joins the sheet
        both = client.put(url, json={**body, "segments": ["tsg", "OE"]})
        assert both.json()["segments"] == ["TSG", "OE"]
        prev = client.get("/api/export/daily-collection/preview",
                          params={"customer_id": world["key"], "from": "2026-08-24",
                                  "to": "2026-08-24"}).json()
        assert "UTR-OE" in {r["Bank Ref"] for r in prev["rows"]}

        bad = client.put(url, json={**body, "recipients": ["not-an-email"]})
        assert bad.status_code == 400
        bad = client.put(url, json={**body, "fiscal_calendar": {"pattern": [4] * 12}})
        assert bad.status_code == 400
    with SessionLocal() as s:
        evs = s.query(AuditLog).filter_by(customer_id=world["pk"],
                                          event_type="config.collection_updated") \
            .order_by(AuditLog.id).all()
        assert len(evs) == 2          # the two saves; the refused ones log nothing
        assert "collection.branch_codes.Metro" in {c["field"] for c in evs[0].details["changes"]}
        assert {"field": "collection.segments", "from": ["TSG"], "to": ["TSG", "OE"]} \
            in evs[1].details["changes"]


def test_ingest_ar_statement(tmp_path):
    from app.main import app
    init_db()
    key = f"dc-ing-{uuid.uuid4().hex[:8]}"
    p = _ar_workbook(tmp_path / "External AR Statement Aug 2026.xlsx", rows=[
        ["CS", "CS-Others", "1331000016", "SR", None, "WABHSR CS -Spares Sales",
         datetime(2026, 7, 1), datetime(2026, 8, 20), "INR", 1000.0, 900.0, None, 5, "Hosur B&S"],
        ["CS", "CS-Others", "1331000017", "SR", None, "WABHSR CS -Spares Sales",
         datetime(2026, 7, 2), datetime(2026, 8, 21), "INR", 50.0, 50.0, None, 5, "Hosur B&S"],
    ])
    pk = None
    try:
        with TestClient(app) as client:
            assert client.post("/api/customers", json={"key": key, "name": "x"}).status_code == 200
            with SessionLocal() as s:
                pk = s.execute(select(Customer.id).where(Customer.key == key)).scalar_one()
            for attempt in range(2):          # the second upload is a replay
                with open(p, "rb") as fh:
                    res = client.post("/api/ingest", data={"customer_id": key},
                                      files={"ar_statement": (p.name, fh)})
                assert res.status_code == 200, res.text
            assert res.json()["files"][0]["outcome"] == "deduped"
        with SessionLocal() as s:
            rows = s.query(GoldArInvoice).filter_by(customer_id=pk).all()
            assert sorted(r.invoice_number for r in rows) == ["1331000016", "1331000017"]
            assert {r.statement_date for r in rows} == {date(2026, 8, 31)}
            rule = s.query(MatchRuleSetRow).filter_by(customer_id=pk).one()
            assert rule.collection_config is None
    finally:
        if pk is not None:
            _wipe(pk, key)
