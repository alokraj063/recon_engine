"""
File storage behind db/storage.py: local disk (default) and S3.

S3 runs against moto's in-memory S3, so no AWS account or network is
needed. The end-to-end tests drive the real API with the S3 backend
swapped in: bronze files land as s3:// references, a reconcile downloads
the statements it needs, and the run workbook is stored in and served
back from S3.
"""

import shutil
import sys
import uuid
from pathlib import Path

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import delete, select

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

boto3 = pytest.importorskip("boto3")
moto = pytest.importorskip("moto")

from db import SessionLocal, init_db  # noqa: E402
from db import storage as storage_mod  # noqa: E402
from db.bronze import register_file  # noqa: E402
from db.ingest import ingest_gold_frames  # noqa: E402
from db.models import (AuditLog, BronzeFile, Customer, ExceptionLedger,  # noqa: E402
                       GoldBankTxn, GoldBill, GoldFileRow, GoldRecovery,
                       MatchLedger, MatchLedgerBill, MatchRuleSetRow, Run,
                       RunFrame, RunMatchBill, SilverRecord, SourceConfig)
from db.storage import (LocalBackend, S3Backend, Storage,  # noqa: E402
                        backend_from_env, is_s3_ref, parse_s3_ref, storage)
from recon.gold import ensure_schema  # noqa: E402

REGION = "ap-south-1"
BRONZE_BUCKET = "test-recon-documents"
RUNS_BUCKET = "test-recon-exports"


# ----------------------------------------------------------------- fixtures
@pytest.fixture
def s3_client(monkeypatch):
    # moto intercepts every call; fake credentials keep botocore from
    # looking for real ones (and from ever reaching AWS)
    for name, value in {"AWS_ACCESS_KEY_ID": "testing",
                        "AWS_SECRET_ACCESS_KEY": "testing",
                        "AWS_SESSION_TOKEN": "testing",
                        "AWS_DEFAULT_REGION": REGION}.items():
        monkeypatch.setenv(name, value)
    monkeypatch.delenv("S3_ENDPOINT_URL", raising=False)
    with moto.mock_aws():
        client = boto3.client("s3", region_name=REGION)
        for bucket in (BRONZE_BUCKET, RUNS_BUCKET):
            client.create_bucket(
                Bucket=bucket,
                CreateBucketConfiguration={"LocationConstraint": REGION})
        yield client


@pytest.fixture
def s3_storage(s3_client):
    """The app-wide `storage` writing to (fake) S3 for one test."""
    previous = storage.backend
    storage.use(S3Backend(BRONZE_BUCKET, RUNS_BUCKET, prefix="recon",
                          client=s3_client))
    try:
        yield storage
    finally:
        storage.use(previous)


def _file(tmp_path, name, content=None):
    p = tmp_path / name
    p.write_bytes(content if content is not None else uuid.uuid4().bytes)
    return p


# ------------------------------------------------------------- references
def test_s3_reference_parsing():
    assert is_s3_ref("s3://b/k/x.pdf") and not is_s3_ref("/data/bronze/x.pdf")
    assert parse_s3_ref("s3://bucket/recon/bronze/c/abc.pdf") == (
        "bucket", "recon/bronze/c/abc.pdf")
    for bad in ("s3://", "s3://bucket", "s3://bucket/", "/local/path"):
        with pytest.raises(ValueError):
            parse_s3_ref(bad)


# ------------------------------------------------------------------- local
def test_local_backend_roundtrip_and_dedup(tmp_path):
    local = Storage(LocalBackend(tmp_path / "root"))
    src = _file(tmp_path, "Statement.PDF", b"pdf-bytes")

    ref = local.save_bronze("cust", src, "abc123")
    assert ref == str(tmp_path / "root" / "bronze" / "cust" / "abc123.pdf")
    assert local.exists(ref)
    # content-addressed: a second save leaves the stored file alone
    Path(ref).write_bytes(b"already stored")
    assert local.save_bronze("cust", src, "abc123") == ref
    assert Path(ref).read_bytes() == b"already stored"

    # local refs are used in place, never copied into the workdir
    workdir = tmp_path / "work"
    assert local.local_path(ref, workdir) == Path(ref)
    assert not workdir.exists()

    wb = local.save_run_workbook("run1", _file(tmp_path, "out.xlsx", b"xlsx"))
    assert Path(wb).read_bytes() == b"xlsx"
    assert local.download_to(wb, tmp_path / "copy.xlsx").read_bytes() == b"xlsx"
    assert not local.exists(str(tmp_path / "missing.xlsx"))
    assert not local.exists(None)


def test_local_prune_keeps_newest(tmp_path):
    import os
    local = Storage(LocalBackend(tmp_path))
    for i in range(5):
        ref = local.save_run_workbook(f"run{i}", _file(tmp_path, f"w{i}.xlsx"))
        os.utime(Path(ref).parent, (1_000_000 + i, 1_000_000 + i))
    local.prune_run_files(keep=2)
    assert sorted(p.name for p in (tmp_path / "runs").iterdir()) == ["run3", "run4"]


# ---------------------------------------------------------------------- s3
def test_s3_bronze_is_content_addressed(s3_client, tmp_path):
    s3 = Storage(S3Backend(BRONZE_BUCKET, RUNS_BUCKET, prefix="/recon/",
                           client=s3_client))
    src = _file(tmp_path, "Bills.XLSX", b"bills-v1")

    ref = s3.save_bronze("wabtec", src, "deadbeef")
    assert ref == f"s3://{BRONZE_BUCKET}/recon/bronze/wabtec/deadbeef.xlsx"
    assert s3.exists(ref)

    # same sha again: no second upload, the stored object is untouched
    src.write_bytes(b"different bytes, same claimed sha")
    assert s3.save_bronze("wabtec", src, "deadbeef") == ref
    body = s3_client.get_object(Bucket=BRONZE_BUCKET,
                                Key="recon/bronze/wabtec/deadbeef.xlsx")["Body"].read()
    assert body == b"bills-v1"


def test_s3_workbook_goes_to_runs_bucket_and_reads_back(s3_client, tmp_path):
    s3 = Storage(S3Backend(BRONZE_BUCKET, RUNS_BUCKET, client=s3_client))
    ref = s3.save_run_workbook("r42", _file(tmp_path, "Recon_Output.xlsx", b"wb"))
    assert ref == f"s3://{RUNS_BUCKET}/runs/r42/Recon_Output.xlsx"
    assert s3.exists(ref)
    assert not s3.exists(f"s3://{RUNS_BUCKET}/runs/nope/Recon_Output.xlsx")

    # download keeps the object's file name (parsers dispatch on suffix)
    local = s3.local_path(ref, tmp_path / "work")
    assert local == tmp_path / "work" / "Recon_Output.xlsx"
    assert local.read_bytes() == b"wb"
    assert s3.download_to(ref, tmp_path / "x.xlsx").read_bytes() == b"wb"

    s3.prune_run_files(keep=0)          # lifecycle rule's job: a no-op
    assert s3.exists(ref)


def test_runs_bucket_defaults_to_bronze_bucket(s3_client, tmp_path):
    s3 = Storage(S3Backend(BRONZE_BUCKET, client=s3_client))
    ref = s3.save_run_workbook("r1", _file(tmp_path, "w.xlsx"))
    assert parse_s3_ref(ref)[0] == BRONZE_BUCKET


def test_s3_explicit_kms_key(s3_client, tmp_path):
    s3 = Storage(S3Backend(BRONZE_BUCKET, client=s3_client,
                           kms_key_id="alias/wabtec-recon-key"))
    ref = s3.save_bronze("c", _file(tmp_path, "s.pdf"), "k1")
    head = s3_client.head_object(Bucket=BRONZE_BUCKET, Key=parse_s3_ref(ref)[1])
    assert head["ServerSideEncryption"] == "aws:kms"


def test_reads_follow_the_reference_not_the_backend(s3_client, tmp_path):
    """A database can hold local refs (older rows) and S3 refs (newer)."""
    s3_ref = Storage(S3Backend(BRONZE_BUCKET, client=s3_client)).save_bronze(
        "c", _file(tmp_path, "a.pdf", b"from-s3"), "s1")
    local_ref = str(_file(tmp_path, "b.pdf", b"from-disk"))

    for writer in (LocalBackend(tmp_path / "root"),
                   S3Backend(BRONZE_BUCKET, client=s3_client)):
        st = Storage(writer)
        assert st.local_path(s3_ref, tmp_path / "w1").read_bytes() == b"from-s3"
        assert st.local_path(local_ref, tmp_path / "w2").read_bytes() == b"from-disk"


def test_backend_from_env(monkeypatch):
    for name in ("STORAGE_BACKEND", "S3_BRONZE_BUCKET", "S3_RUNS_BUCKET",
                 "S3_PREFIX", "S3_KMS_KEY_ID"):
        monkeypatch.delenv(name, raising=False)
    assert isinstance(backend_from_env(), LocalBackend)

    monkeypatch.setenv("STORAGE_BACKEND", "S3")
    with pytest.raises(ValueError, match="S3_BRONZE_BUCKET"):
        backend_from_env()
    monkeypatch.setenv("S3_BRONZE_BUCKET", "docs")
    monkeypatch.setenv("S3_PREFIX", "recon")
    b = backend_from_env()
    assert (b.bronze_bucket, b.runs_bucket, b.prefix) == ("docs", "docs", "recon/")

    monkeypatch.setenv("STORAGE_BACKEND", "gcs")
    with pytest.raises(ValueError, match="STORAGE_BACKEND"):
        backend_from_env()


# ------------------------------------------------------ end to end via API
@pytest.fixture
def world(s3_storage, tmp_path):
    """Throwaway customer whose bronze files are registered while S3 is the
    storage backend."""
    init_db()
    from fastapi import FastAPI
    from app.routes import router
    app = FastAPI()
    app.include_router(router)
    client = TestClient(app)
    key = f"s3store-{uuid.uuid4().hex[:8]}"
    r = client.post("/api/customers", json={"key": key, "name": "s3 storage test"})
    assert r.status_code == 200, r.text
    with SessionLocal() as s:
        c = s.execute(select(Customer).where(Customer.key == key)).scalar_one()
        bz = {}
        for slot, name, source_type in (("stmt", "stmt-a.txt", "bank_statement"),
                                        ("bills", "bills.txt", "bill_status")):
            bz[slot] = register_file(s, c, source_type,
                                     _file(tmp_path, name), name).id
        s.commit()
        pk = c.id
    yield {"pk": pk, "key": key, "bz": bz, "client": client}
    with SessionLocal() as s:
        run_ids = select(Run.id).where(Run.customer_id == pk)
        match_ids = select(MatchLedger.id).where(MatchLedger.customer_id == pk)
        s.execute(delete(RunMatchBill).where(RunMatchBill.run_id.in_(run_ids)))
        s.execute(delete(RunFrame).where(RunFrame.run_id.in_(run_ids)))
        s.execute(delete(MatchLedgerBill).where(MatchLedgerBill.match_ledger_id.in_(match_ids)))
        for model in (MatchLedger, ExceptionLedger, AuditLog, Run, GoldFileRow,
                      GoldRecovery, GoldBill, GoldBankTxn, SilverRecord, BronzeFile,
                      SourceConfig, MatchRuleSetRow):
            s.execute(delete(model).where(model.customer_id == pk))
        s.execute(delete(Customer).where(Customer.id == pk))
        s.commit()
    shutil.rmtree(storage.root / "bronze" / key, ignore_errors=True)


def _load_matching_pair(pk, bz):
    credit = ensure_schema(pd.DataFrame([{
        "bank_ref": "T1", "value_date": pd.Timestamp("2026-03-18"),
        "amount": 5000.0, "used_in_recon": True,
        # a zone makes it IREPS money: a credit with no match signal is a
        # non-IREPS receipt and an incremental run never matches it
        "zone_guess": "NR"}]), "bank_txns")
    credit["row_seq"] = range(len(credit))
    bill = ensure_schema(pd.DataFrame([{
        "bill_number": "INV-1", "submission_ref": "CO6-1",
        "submission_date": pd.Timestamp("2026-03-10"),
        "net_payable_amount": 5000.0, "bill_status": "PAYMENT MADE",
        "payment_advice_date": pd.Timestamp("2026-03-18"),
        "payment_order_ref": "CO7-1",
        "payment_order_date": pd.Timestamp("2026-03-17"), "zone": "NR",
        "data_row": 2}]), "bills")
    bill["row_seq"] = range(len(bill))
    with SessionLocal() as s:
        ingest_gold_frames(s, pk, {"bank_txns": credit}, {"bank_txns": bz["stmt"]})
        ingest_gold_frames(s, pk, {"bills": bill}, {"bills": bz["bills"]})
        s.commit()


@pytest.mark.parametrize("mode", ["snapshot", "incremental"])
def test_reconcile_and_workbook_download_on_s3(world, s3_client, monkeypatch, mode):
    pk, bz, client, key = world["pk"], world["bz"], world["client"], world["key"]

    with SessionLocal() as s:
        refs = {slot: s.get(BronzeFile, bid).stored_path for slot, bid in bz.items()}
    for ref in refs.values():
        bucket, obj_key = parse_s3_ref(ref)
        assert bucket == BRONZE_BUCKET
        assert obj_key.startswith(f"recon/bronze/{key}/")
        s3_client.head_object(Bucket=bucket, Key=obj_key)

    downloads = []
    real_local_path = storage_mod.Storage.local_path

    def spy(self, ref, workdir):
        path = real_local_path(self, ref, workdir)
        downloads.append((ref, path, path.exists()))
        return path

    monkeypatch.setattr(storage_mod.Storage, "local_path", spy)

    _load_matching_pair(pk, bz)
    r = client.post("/api/reconcile", json={
        "customer_id": key, "statement_bronze_ids": [bz["stmt"]], "mode": mode})
    assert r.status_code == 200, r.text
    run_id = r.json()["run_id"]

    # the statement was fetched from S3 to a real file for the selfcheck,
    # and that request-scoped temp copy is gone once the request finished
    assert [(ref, existed) for ref, _, existed in downloads] == [(refs["stmt"], True)]
    assert not downloads[0][1].exists()

    with SessionLocal() as s:
        wb_ref = s.get(Run, run_id).workbook_path
    assert wb_ref == f"s3://{RUNS_BUCKET}/recon/runs/{run_id}/Recon_Output.xlsx"
    stored_bytes = s3_client.get_object(
        Bucket=RUNS_BUCKET, Key=parse_s3_ref(wb_ref)[1])["Body"].read()

    dl = client.get(f"/api/runs/{run_id}/workbook")
    assert dl.status_code == 200, dl.text
    assert dl.content[:2] == b"PK"                      # an xlsx (zip)
    if mode == "snapshot":
        # nothing in the ledger: the stored object streams back unchanged
        assert dl.content == stored_bytes
    else:
        # incremental: ledger sheets appended at download time
        from io import BytesIO
        from openpyxl import load_workbook
        sheets = load_workbook(BytesIO(dl.content), read_only=True).sheetnames
        assert "Decisions" in sheets

    # a run whose workbook object is missing is a 404, not a 500
    s3_client.delete_object(Bucket=RUNS_BUCKET, Key=parse_s3_ref(wb_ref)[1])
    assert client.get(f"/api/runs/{run_id}/workbook").status_code == 404
