"""
db/webadi_export.py + GET /api/export/webadi: confirmed matches -> Oracle
AR Receipt Upload records, one per (credit, bill, recovery line), every
field repeated across a bill's recovery rows; blanks where the data
cannot say (unmapped head, unknown zone, Friction's method).
"""

import sys
import uuid
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from db import SessionLocal, init_db  # noqa: E402
from db import webadi_export as w  # noqa: E402
from db.bronze import register_file  # noqa: E402
from db.models import (AuditLog, BronzeFile, Customer, GoldBankTxn,  # noqa: E402
                       GoldBill, GoldRecovery, MatchLedger, MatchLedgerBill)
from db.storage import storage  # noqa: E402

DAY = date(2025, 10, 10)


@pytest.fixture()
def world(tmp_path):
    init_db()
    key = f"webadi-{uuid.uuid4().hex[:8]}"
    with SessionLocal() as s:
        c = Customer(key=key, name="webadi test")
        s.add(c)
        s.flush()
        p = tmp_path / "f.txt"
        p.write_text(uuid.uuid4().hex)
        fid = register_file(s, c, "bill_status", p, p.name).id
        seq = iter(range(1000))

        def bill(number, unit, zone, net, heads=()):
            b = GoldBill(customer_id=c.id, bronze_file_id=fid, row_seq=next(seq),
                         bill_number=number, submission_ref=f"CO6-{number}",
                         zone=zone, operating_unit=unit, net_payable_amount=net,
                         deduction_amount=sum(a for _h, a in heads),
                         bill_status="PAYMENT MADE")
            s.add(b)
            s.flush()
            for head, amt in heads:
                s.add(GoldRecovery(customer_id=c.id, bronze_file_id=fid,
                                   row_seq=next(seq), gold_bill_id=b.id,
                                   recovery_head=head, recovery_amt=amt))
            return b

        def match(ref, amount, bills, status="LOCKED", day=DAY):
            t = GoldBankTxn(customer_id=c.id, bronze_file_id=fid, row_seq=next(seq),
                            bank_ref=ref, amount=amount, value_date=day,
                            used_in_recon=True)
            s.add(t)
            s.flush()
            m = MatchLedger(customer_id=c.id, match_id="m", gold_bank_txn_id=t.id,
                            confidence="HIGH", status=status, seq=next(seq))
            s.add(m)
            s.flush()
            for b in bills:
                s.add(MatchLedgerBill(match_ledger_id=m.id, gold_bill_id=b.id,
                                      role="picked"))
            return m

        # Hosur ICF bill with three recoveries, one unmapped
        b1 = bill("1331000016", "Hosur", "ICF", 1176.0, [
            ("AMT RECOVERED-LIQUIDATED DAMAGES", 100.0),
            ("GST TDS DEDUCTION", 20.0),
            ("GROUND RENT", 5.0)])
        match("UTR001", 1176.0, [b1])
        # Rohtak batch: two bills, one with no recoveries, one OE security deposit
        b2 = bill("R-1", "Rohtak", "NFR", 500.0)
        b3 = bill("R-2", "Rohtak", "CLW", 300.0, [("DEPOSIT SECURITY", 7.0)])
        match("UTR002", 800.0, [b2, b3])
        # Friction bill, unknown zone
        b4 = bill("F-1", "Friction", "MTPK", 50.0, [("INCOME TAX - CONTR (Company)", 1.0)])
        match("UTR003", 50.0, [b4])
        # excluded: an OPEN review match, and a LOCKED one on another day
        match("UTR-OPEN", 10.0, [bill("X-1", "Hosur", "CR", 10.0)], status="OPEN")
        match("UTR-LATER", 20.0, [bill("X-2", "Hosur", "CR", 20.0)], day=date(2025, 10, 11))
        s.commit()
        pk = c.id
    yield {"pk": pk, "key": key}
    with SessionLocal() as s:
        ids = [m.id for m in s.query(MatchLedger).filter_by(customer_id=pk)]
        s.execute(delete(MatchLedgerBill).where(MatchLedgerBill.match_ledger_id.in_(ids)))
        for model in (MatchLedger, AuditLog, GoldRecovery, GoldBill, GoldBankTxn, BronzeFile):
            s.execute(delete(model).where(model.customer_id == pk))
        s.execute(delete(Customer).where(Customer.id == pk))
        s.commit()
    import shutil
    shutil.rmtree(storage.root / "bronze" / key, ignore_errors=True)


def test_records(world):
    with SessionLocal() as s:
        recs = w.build_records(s, world["pk"], DAY, DAY)
    assert {r["Receipt Number"] for r in recs} == {"UTR001", "UTR002", "UTR003"}

    hosur = [r for r in recs if r["Receipt Number"] == "UTR001"]
    assert len(hosur) == 3          # one row per recovery line
    for r in hosur:                 # every other field repeated
        assert r["Operating Unit ID Selected"] == w.OU_NAMES["Hosur"]
        assert r["Customer Name"] == "ICF-INTEGRAL COACH FACTORY"
        assert r["Receipt Method"] == "WABHSR_Receipt_HSBC_INR-7001"
        assert r["Receipt Date"] == "10-Oct-2025"
        assert r["Invoice Number"] == "1331000016"
        assert r["Receipt Amount Applied"] == 1176.0
        assert r["Currency"] == "INR" and r["Upl"] == "O"
    assert [(r["Adjustment Amount"], r["Adjustment Type"]) for r in hosur] == [
        (100.0, "WABHSR - LD"), (20.0, "WABHSR GST - TDS - IGST"), (5.0, None)]
    assert hosur[2]["issues"] == ["NO_ADJUSTMENT_TYPE"]

    rohtak = [r for r in recs if r["Receipt Number"] == "UTR002"]
    assert [(r["Invoice Number"], r["Adjustment Type"]) for r in rohtak] == [
        ("R-1", None), ("R-2", "WABRTK Security Deposit - OE")]
    assert all(r["Receipt Amount"] == 800.0 for r in rohtak)
    assert rohtak[0]["Customer Name"] == "NEFR-NORTHEAST FRONTIER RAILWAY"
    assert rohtak[0]["issues"] == []

    friction = [r for r in recs if r["Receipt Number"] == "UTR003"][0]
    assert friction["Receipt Method"] is None and friction["Adjustment Type"] is None
    assert friction["Customer Name"] is None
    assert set(friction["issues"]) == {"NO_CUSTOMER", "NO_RECEIPT_METHOD",
                                       "NO_ADJUSTMENT_TYPE"}


def test_adjustment_mapping():
    assert w.adjustment_type("DEPOSIT STORES-LIQUIDITY DAMAGE", "Hosur", "TSG") == "WABHSR - LD"
    assert w.adjustment_type("INCOME TAX - CONTR(Section-194Q)", "Hosur", "TSG") == "WABHSR IT TDS"
    assert w.adjustment_type("DEPOSIT SECURITY", "Hosur", "TSG") == "WABHSR Security Deposit - CS"
    assert w.adjustment_type("PENALTY CHARGES", "Hosur", "OE") == "WABHSR General Damages"
    assert w.adjustment_type("OTHER CHARGES", "Rohtak", None) == "WABRTK IR other Deductions"
    assert w.adjustment_type("DEPOSIT EXP", "Hosur", "TSG") is None


def test_download(world):
    from app.main import app
    with TestClient(app) as client:
        res = client.get("/api/export/webadi",
                         params={"customer_id": world["key"], "from": "2025-10-10"})
        assert res.status_code == 200
        wb = load_workbook(BytesIO(res.content))
        assert wb.sheetnames == ["Hosur", "Rohtak", "Friction"]
        ws = wb["Hosur"]
        assert ws["D5"].value == "Receivables Super User GUI - WABHSR"
        assert ws["B15"].value == "Upl" and ws["R15"].value == "Messages"
        assert [ws.cell(row=r, column=6).value for r in (17, 18, 19)] == ["UTR001"] * 3
        assert ws["B20"].value is None

        # the unit filter narrows the download; preview agrees
        prev = client.get("/api/export/webadi/preview",
                          params={"customer_id": world["key"], "from": "2025-10-10",
                                  "unit": "Rohtak"}).json()
        assert prev["summary"]["records"] == 2 and prev["summary"]["receipts"] == 1
    with SessionLocal() as s:
        ev = s.query(AuditLog).filter_by(customer_id=world["pk"],
                                         event_type="export.webadi_downloaded").one()
        assert ev.details["records"] == 6


def test_gl_date_and_deduction_gap(world):
    with SessionLocal() as s:
        # GL Date defaults to the receipt date; an explicit one wins
        recs = w.build_records(s, world["pk"], DAY, DAY)
        assert {r["GL Date"] for r in recs} == {"10-Oct-2025"}
        recs = w.build_records(s, world["pk"], DAY, DAY, gl_date=date(2025, 10, 12))
        assert {r["GL Date"] for r in recs} == {"12-Oct-2025"}
        # a deduction the recovery lines do not cover is flagged, never silent
        b = s.query(GoldBill).filter_by(customer_id=world["pk"], bill_number="R-1").one()
        b.deduction_amount = 12.0
        s.flush()
        r1 = [r for r in w.build_records(s, world["pk"], DAY, DAY)
              if r["Invoice Number"] == "R-1"]
        assert r1[0]["issues"] == ["DEDUCTION_WITHOUT_DETAIL"]
        s.rollback()
